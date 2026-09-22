"""Deterministic ING-02a golden + hostile fixture generator (INGEST-GOLDEN-001).

Regenerates the committed fixtures under ``tests/fixtures/sources/`` and their
canonical goldens under ``goldens/`` with a FIXED source-version UUID, so the
golden JSON is byte-stable for a given fixture content and lockfile-pinned
parser versions.

Binary office fixtures are not byte-reproducible (the office libraries embed
zip timestamps), so regenerate goldens after regenerating binaries:
``uv run python tests/fixtures/sources/generate_fixtures.py``.

The hostile corpus is intentionally malicious:
* ``hostile-macro.xlsx``  — xl/vbaProject.bin member (macros; policy-blocked)
* ``hostile-extref.docx`` — external non-hyperlink relationship (policy-blocked)
* ``hostile-bomb.xlsx``   — central directory declares a 2 GB member (too_large)
* ``hostile-corrupt.xlsx``— zip magic with garbage payload (corrupt)
"""

from __future__ import annotations

import io
import json
import re
import struct
import sys
import uuid
import zipfile
from pathlib import Path

import docx
import openpyxl
import pptx
from ebooklib import epub
from milpbooklm_adapters.parsers.isolation import (
    IsolatedParser,
    ParseSuccess,
    WebLocatorContext,
)

FIXTURES = Path(__file__).resolve().parent
GOLDENS = FIXTURES / "goldens"
SOURCE_VERSION_ID = uuid.UUID("00000000-0000-5000-8000-000000000211")

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
HTML_TYPE = "text/html"
EPUB_TYPE = "application/epub+zip"
WEB_CONTEXT = WebLocatorContext(
    canonical_url="https://example.test/reports/q3",
    captured_at="2026-09-22T12:00:00+00:00",
)


def main() -> int:
    if sys.argv[1:] == ["t22"]:
        _write_t22_fixtures()
        _regenerate_t22_goldens()
        print("task 22 fixtures + goldens written")
        return 0
    _write("golden-markdown.md", _markdown_fixture())
    _write("golden-csv.csv", _csv_fixture())
    _write("golden-csv-utf16le.csv", _utf16_csv_fixture())
    _write("golden-xlsx.xlsx", _xlsx_fixture())
    _write("golden-docx.docx", _docx_fixture())
    _write("golden-pptx.pptx", _pptx_fixture())
    _write("hostile-macro.xlsx", _macro_xlsx_fixture())
    _write("hostile-extref.docx", _extref_docx_fixture())
    _write("hostile-bomb.xlsx", _bomb_xlsx_fixture())
    _write("hostile-corrupt.xlsx", _corrupt_fixture())
    _regenerate_goldens()
    print("fixtures + goldens written")
    return 0


def _write_t22_fixtures() -> None:
    epub_data = _epub_fixture()
    _write("golden-html.html", _html_fixture())
    _write("golden-epub.epub", epub_data)
    _write("hostile-bomb.epub", _forge_member_size(epub_data, _epub_chapter_name(epub_data), 2**31))
    _write("hostile-corrupt.epub", _corrupt_fixture())
    _write("hostile-entity.epub", _entity_epub_fixture(epub_data))
    _write("hostile-external.epub", _external_epub_fixture(epub_data))


def _write(name: str, data: bytes) -> None:
    (FIXTURES / name).write_bytes(data)


def _markdown_fixture() -> bytes:
    body = (
        "\ufeff# Vierteljahresbericht\r\n\r\n"
        "Eine Einleitung mit **Hervorhebung** und einem [Link](https://example.org).\r\n\r\n"
        "## Zahlen\r\n\r\n"
        "- Umsatz stieg\r\n- Kosten sanken\r\n\r\n"
        "1. Erstens\r\n2. Zweitens\r\n\r\n"
        "> Eine zitierte Aussage aus dem Bericht.\r\n\r\n"
        "```python\r\nsumme = sum(range(10))\r\n```\r\n\r\n"
        "| Kennzahl | 2025 | 2026 |\r\n|---|---:|---:|\r\n| Umsatz | 120 | 180 |\r\n"
    )
    return body.encode("utf-8")


def _html_fixture() -> bytes:
    return (
        "<!doctype html><html><head><title>Atlas</title></head><body>"
        "<nav>Navigationstext</nav><main><h1>Quartalsbericht</h1>"
        "<p>Der Umsatz stieg.</p><h2>Ausblick</h2><p>Das Team wächst.</p>"
        "</main><script>prompt injection from script</script></body></html>"
    ).encode()


def _epub_fixture() -> bytes:
    book = epub.EpubBook()
    book.set_identifier("milpbooklm-t22")
    book.set_title("Atlas EPUB")
    book.set_language("de")
    chapter = epub.EpubHtml(title="Quartalsbericht", file_name="chapter.xhtml", lang="de")
    chapter.content = "<h1>Quartalsbericht</h1><p>Der Umsatz stieg.</p>"
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    buffer = io.BytesIO()
    epub.write_epub(buffer, book)
    return buffer.getvalue()


def _epub_chapter_name(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return next(name for name in archive.namelist() if name.endswith("chapter.xhtml"))


def _entity_epub_fixture(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = {info.filename: archive.read(info) for info in archive.infolist()}
    chapter_name = _epub_chapter_name(data)
    members[chapter_name] = (
        b'<!DOCTYPE html [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
        + members[chapter_name]
    )
    return _rezip(members)


def _external_epub_fixture(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = {info.filename: archive.read(info) for info in archive.infolist()}
    chapter_name = _epub_chapter_name(data)
    members[chapter_name] = members[chapter_name].replace(
        b"</body>", b'<img src="https://attacker.invalid/pixel" /></body>'
    )
    return _rezip(members)


def _csv_fixture() -> bytes:
    body = (
        "Name;Stadt;Bemerkung\r\n"
        'Anna;Berlin;"Halbtags; Montag"\r\n'
        "Ben;Wien;Zuständig für Einkauf\r\n"
        "Clara;Zürich;\r\n"
    )
    return body.encode("utf-8")


def _utf16_csv_fixture() -> bytes:
    body = "Name,City,Notes\r\nAnna,Berlin,\"Team A\"\r\n"
    return b"\xff\xfe" + body.encode("utf-16-le")


def _xlsx_fixture() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Daten"
    sheet["A1"] = "Kennzahl"
    sheet["B1"] = "Wert"
    sheet["A2"] = "Umsatz"
    sheet["B2"] = 120
    sheet["A3"] = "Kosten"
    sheet["B3"] = 80
    sheet["A4"] = "Saldo"
    sheet["B4"] = "=B2-B3"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _with_cached_formula_value(buffer.getvalue(), "B4", "40")


def _with_cached_formula_value(data: bytes, cell_ref: str, value: str) -> bytes:
    """Inject a cached <v> result into one formula cell of the first sheet."""
    archive = zipfile.ZipFile(io.BytesIO(data))
    members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    archive.close()
    sheet_name = "xl/worksheets/sheet1.xml"
    sheet = members[sheet_name].decode("utf-8")
    pattern = re.compile(rf'(<c r="{cell_ref}"[^>]*>)(<f>[^<]*</f>)(?:<v>[^<]*</v>)?(</c>)')
    patched, count = pattern.subn(rf"\g<1>\g<2><v>{value}</v>\g<3>", sheet)
    if count != 1:
        raise SystemExit(f"fixture generation failed: formula cell {cell_ref} not found")
    members[sheet_name] = patched.encode("utf-8")
    return _rezip(members)


def _rezip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def _docx_fixture() -> bytes:
    document = docx.Document()
    document.add_heading("Vierteljahresbericht Q3", level=1)
    document.add_paragraph("Der Umsatz stieg im Berichtszeitraum deutlich.")
    document.add_heading("Personal", level=2)
    document.add_paragraph("Das Team wuchs auf zwölf Personen.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kennzahl"
    table.cell(0, 1).text = "Wert"
    table.cell(1, 0).text = "Umsatz"
    table.cell(1, 1).text = "120"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_fixture() -> bytes:
    presentation = pptx.Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = "Vierteljahresbericht"
    first.placeholders[1].text = "Ergebnisse des dritten Quartals"
    first.notes_slide.notes_text_frame.text = "Sprechernotiz: Umsatz hervorheben."
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Ausblick"
    body = second.placeholders[1]
    body.text = "Erstes Halbjahr"
    paragraph = body.text_frame.add_paragraph()
    paragraph.text = "Zweites Halbjahr"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _macro_xlsx_fixture() -> bytes:
    base = _xlsx_fixture()
    archive = zipfile.ZipFile(io.BytesIO(base))
    members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    archive.close()
    members["xl/vbaProject.bin"] = b"\x00\x01\x02MACROPAYLOAD-STUB"
    return _rezip(members)


def _extref_docx_fixture() -> bytes:
    base = _docx_fixture()
    archive = zipfile.ZipFile(io.BytesIO(base))
    members = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    archive.close()
    rels_name = "word/_rels/document.xml.rels"
    rels = members[rels_name].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdHostile"'
        ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
        ' Target="file://\\\\host\\share\\logo.png" TargetMode="External"'
        " /></Relationships>",
    )
    members[rels_name] = rels.encode("utf-8")
    return _rezip(members)


def _bomb_xlsx_fixture() -> bytes:
    base = _xlsx_fixture()
    return _forge_member_size(base, "xl/worksheets/sheet1.xml", 2 * 1024 * 1024 * 1024)


def _corrupt_fixture() -> bytes:
    return b"PK\x03\x04" + b"\xde\xad\xbe\xef" * 64


def _forge_member_size(data: bytes, member: str, size: int) -> bytes:
    """Overwrite one central-directory uncompressed-size field (nominal ratio bomb)."""
    raw = bytearray(data)
    offset = 0
    while True:
        offset = raw.find(b"PK\x01\x02", offset)
        if offset < 0:
            raise SystemExit(f"fixture generation failed: central record for {member} not found")
        name_length = struct.unpack_from("<H", raw, offset + 28)[0]
        extra_length = struct.unpack_from("<H", raw, offset + 30)[0]
        comment_length = struct.unpack_from("<H", raw, offset + 32)[0]
        name = bytes(raw[offset + 46 : offset + 46 + name_length]).decode("utf-8")
        if name == member:
            struct.pack_into("<I", raw, offset + 24, size)
            return bytes(raw)
        offset += 46 + name_length + extra_length + comment_length


def _regenerate_goldens() -> None:
    parser = IsolatedParser()
    media = {
        "golden-markdown.md": "text/markdown",
        "golden-csv.csv": "text/csv",
        "golden-csv-utf16le.csv": "text/csv",
        "golden-xlsx.xlsx": XLSX_TYPE,
        "golden-docx.docx": DOCX_TYPE,
        "golden-pptx.pptx": PPTX_TYPE,
    }
    GOLDENS.mkdir(parents=True, exist_ok=True)
    for name, media_type in media.items():
        result = parser.parse(
            SOURCE_VERSION_ID, media_type, (FIXTURES / name).read_bytes()
        )
        if not isinstance(result, ParseSuccess):
            raise SystemExit(f"golden regeneration failed for {name}: {result}")
        golden_path = GOLDENS / f"{Path(name).stem}.canonical.json"
        golden_path.write_text(json.dumps(result.document.to_json(), indent=1) + "\n")


def _regenerate_t22_goldens() -> None:
    parser = IsolatedParser()
    inputs = {
        "golden-html.html": (HTML_TYPE, WEB_CONTEXT),
        "golden-epub.epub": (EPUB_TYPE, None),
    }
    GOLDENS.mkdir(parents=True, exist_ok=True)
    for name, (media_type, context) in inputs.items():
        result = parser.parse(
            SOURCE_VERSION_ID,
            media_type,
            (FIXTURES / name).read_bytes(),
            web_locator=context,
        )
        if not isinstance(result, ParseSuccess):
            raise SystemExit(f"golden regeneration failed for {name}: {result}")
        golden_path = GOLDENS / f"{Path(name).stem}.canonical.json"
        golden_path.write_text(json.dumps(result.document.to_json(), indent=1) + "\n")


if __name__ == "__main__":
    raise SystemExit(main())
