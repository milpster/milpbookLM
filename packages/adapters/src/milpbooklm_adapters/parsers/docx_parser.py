"""
Deterministic DOCX canonical parser (python-docx structural parse).

Body blocks (paragraphs and tables) are walked in document order; heading
styles become heading nodes with their level, tables become row/cell trees.
Macros and non-hyperlink external references are policy-rejected before
python-docx loads anything. Locator guarantees: structural block/row/cell
numbering — DOCX stores no character offsets, so none are invented.
"""

from __future__ import annotations

import io
import re
import uuid
from importlib.metadata import version
from typing import Any, Final

from milpbooklm_contracts.canonical_document import (
    CanonicalDocument,
    CanonicalNode,
    JsonValue,
    NodeKind,
    ParserDescriptor,
    SourceLocator,
    canonical_document_id,
    canonical_node_id,
)

from .office import (
    OfficeCorruptError,
    assert_no_external_references,
    assert_no_macros,
    open_bounded_archive,
)

_PROFILE: Final = "docx-structure-v1"
_HEADING_STYLE: Final = re.compile(r"^heading\s*(\d+)?$")


def parse_docx(source_version_id: uuid.UUID, data: bytes) -> CanonicalDocument:
    """Extract paragraphs, headings, and tables in document order."""
    archive = open_bounded_archive(data)
    assert_no_macros(archive, "word/")
    assert_no_external_references(archive)
    archive.close()
    import docx  # noqa: PLC0415 - parser imports load only inside the child

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise OfficeCorruptError("document cannot be opened by python-docx") from exc
    parser = ParserDescriptor(
        identity="milpbooklm.docx",
        version="1",
        profile=_PROFILE,
        tool_versions=(f"python-docx={version('python-docx')}",),
    )
    document_id = canonical_document_id(source_version_id, parser)
    root_id = canonical_node_id(document_id, "document")
    nodes = [
        CanonicalNode(
            node_id=root_id,
            kind=NodeKind.DOCUMENT,
            structural_identity="document",
            parent_id=None,
            child_order=0,
            text=None,
            locator=SourceLocator(path=("document",)),
        )
    ]
    for block_number, block in enumerate(_body_blocks(document), start=1):
        if block.get("kind") == "paragraph":
            _emit_paragraph(document_id, nodes, root_id, block_number, block["value"])
        else:
            _emit_table(document_id, nodes, root_id, block_number, block["value"])
    return CanonicalDocument(
        document_id=document_id,
        source_version_id=source_version_id,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        parser=parser,
        languages=(),
        metadata={},
        root_node_ids=(root_id,),
        nodes=tuple(nodes),
    )


def _body_blocks(document: Any) -> list[dict[str, object]]:
    """Yield paragraphs and tables interleaved in true document order."""
    import docx.table  # noqa: PLC0415
    import docx.text.paragraph  # noqa: PLC0415
    from docx.oxml.table import CT_Tbl  # noqa: PLC0415
    from docx.oxml.text.paragraph import CT_P  # noqa: PLC0415

    blocks: list[dict[str, object]] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            blocks.append(
                {"kind": "paragraph", "value": docx.text.paragraph.Paragraph(child, document)}
            )
        elif isinstance(child, CT_Tbl):
            blocks.append({"kind": "table", "value": docx.table.Table(child, document)})
    return blocks


def _heading_level(paragraph: Any) -> int | None:
    try:
        name = paragraph.style.name or ""
    except Exception:
        return None
    match = _HEADING_STYLE.match(name.strip().lower())
    if match is None or match.group(1) is None:
        return 1 if match else None
    return min(int(match.group(1)), 6)


def _emit_paragraph(
    document_id: uuid.UUID,
    nodes: list[CanonicalNode],
    root_id: uuid.UUID,
    block_number: int,
    paragraph: Any,
) -> None:
    text = paragraph.text
    level = _heading_level(paragraph)
    kind = NodeKind.HEADING if level is not None else NodeKind.PARAGRAPH
    identity = f"block:{block_number}"
    extra: dict[str, JsonValue] = {}
    if level is not None:
        extra["heading_level"] = level
    nodes.append(
        CanonicalNode(
            node_id=canonical_node_id(document_id, identity),
            kind=kind,
            structural_identity=identity,
            parent_id=root_id,
            child_order=block_number,
            text=text,
            locator=SourceLocator(
                path=("document", "block", str(block_number)),
                extra_fields={"block": block_number},
            ),
            extra_fields=extra,
        )
    )


def _emit_table(
    document_id: uuid.UUID,
    nodes: list[CanonicalNode],
    root_id: uuid.UUID,
    block_number: int,
    table: Any,
) -> None:
    table_identity = f"block:{block_number}:table"
    table_id = canonical_node_id(document_id, table_identity)
    nodes.append(
        CanonicalNode(
            node_id=table_id,
            kind=NodeKind.TABLE,
            structural_identity=table_identity,
            parent_id=root_id,
            child_order=block_number,
            text=None,
            locator=SourceLocator(
                path=("document", "block", str(block_number), "table"),
                extra_fields={"block": block_number},
            ),
        )
    )
    for row_number, row in enumerate(table.rows, start=1):
        row_identity = f"{table_identity}:row:{row_number}"
        row_id = canonical_node_id(document_id, row_identity)
        nodes.append(
            CanonicalNode(
                node_id=row_id,
                kind=NodeKind.ROW,
                structural_identity=row_identity,
                parent_id=table_id,
                child_order=row_number,
                text=None,
                locator=SourceLocator(
                    path=("document", "block", str(block_number), "row", str(row_number)),
                    extra_fields={"block": block_number, "row": row_number},
                ),
            )
        )
        for column_number, cell in enumerate(row.cells, start=1):
            cell_identity = f"{row_identity}:cell:{column_number}"
            cell_text = "\n".join(
                paragraph.text for paragraph in cell.paragraphs if paragraph.text
            )
            nodes.append(
                CanonicalNode(
                    node_id=canonical_node_id(document_id, cell_identity),
                    kind=NodeKind.CELL,
                    structural_identity=cell_identity,
                    parent_id=row_id,
                    child_order=column_number,
                    text=cell_text,
                    locator=SourceLocator(
                        path=(
                            "document",
                            "block",
                            str(block_number),
                            "row",
                            str(row_number),
                        ),
                        extra_fields={
                            "block": block_number,
                            "row": row_number,
                            "col": column_number,
                        },
                    ),
                )
            )
