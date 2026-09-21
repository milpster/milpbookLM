"""
Deterministic PPTX canonical parser (python-pptx).

Slide and shape locators are structural: slide number, shape index, and
paragraph index within the shape's text frame. Speaker notes are a separate
stream — nodes whose structural identity carries the ``notes`` segment
under the slide, never merged into shape text. Macros and non-hyperlink
external references are policy-rejected before python-pptx loads anything.
"""

from __future__ import annotations

import io
import uuid
from importlib.metadata import version
from typing import Any, Final

from milpbooklm_contracts.canonical_document import (
    CanonicalDocument,
    CanonicalNode,
    JsonValue,
    NodeKind,
    ParserDescriptor,
    SourceLocator,
    canonical_document_id,
    canonical_node_id,
)

from .office import (
    OfficeCorruptError,
    OfficeTooLargeError,
    assert_no_external_references,
    assert_no_macros,
    open_bounded_archive,
)

_PROFILE: Final = "pptx-slides-v1"
_MAX_SLIDES: Final = 500
_MAX_SHAPES: Final = 512
_MAX_PARAGRAPHS: Final = 2_048


def parse_pptx(source_version_id: uuid.UUID, data: bytes) -> CanonicalDocument:
    """Extract slide shape text and separate speaker notes."""
    archive = open_bounded_archive(data)
    assert_no_macros(archive, "ppt/")
    assert_no_external_references(archive)
    archive.close()
    import pptx  # noqa: PLC0415 - parser imports load only inside the child

    try:
        presentation = pptx.Presentation(io.BytesIO(data))
    except Exception as exc:
        raise OfficeCorruptError("presentation cannot be opened by python-pptx") from exc
    if len(presentation.slides) > _MAX_SLIDES:
        raise OfficeTooLargeError("presentation exceeds slide limit")
    parser = ParserDescriptor(
        identity="milpbooklm.pptx",
        version="1",
        profile=_PROFILE,
        tool_versions=(f"python-pptx={version('python-pptx')}",),
    )
    document_id = canonical_document_id(source_version_id, parser)
    root_id = canonical_node_id(document_id, "document")
    nodes = [
        CanonicalNode(
            node_id=root_id,
            kind=NodeKind.DOCUMENT,
            structural_identity="document",
            parent_id=None,
            child_order=0,
            text=None,
            locator=SourceLocator(path=("document",)),
        )
    ]
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_identity = f"slide:{slide_number}"
        slide_id = canonical_node_id(document_id, slide_identity)
        nodes.append(
            CanonicalNode(
                node_id=slide_id,
                kind=NodeKind.SLIDE,
                structural_identity=slide_identity,
                parent_id=root_id,
                child_order=slide_number,
                text=None,
                locator=SourceLocator(
                    path=("document", "slide", str(slide_number)),
                    extra_fields={"slide": slide_number},
                ),
            )
        )
        _emit_shapes(document_id, nodes, slide_id, slide_number, slide)
        _emit_notes(document_id, nodes, slide_id, slide_number, slide)
    return CanonicalDocument(
        document_id=document_id,
        source_version_id=source_version_id,
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        parser=parser,
        languages=(),
        metadata={"slides": len(presentation.slides)},
        root_node_ids=(root_id,),
        nodes=tuple(nodes),
    )


def _emit_shapes(
    document_id: uuid.UUID,
    nodes: list[CanonicalNode],
    slide_id: uuid.UUID,
    slide_number: int,
    slide: Any,
) -> None:
    shapes = list(slide.shapes)
    if len(shapes) > _MAX_SHAPES:
        raise OfficeTooLargeError("slide exceeds shape limit")
    for shape_number, shape in enumerate(shapes, start=1):
        if not shape.has_text_frame:
            continue
        is_title = _is_title_placeholder(shape)
        paragraphs = list(shape.text_frame.paragraphs)
        if len(paragraphs) > _MAX_PARAGRAPHS:
            raise OfficeTooLargeError("shape exceeds paragraph limit")
        for paragraph_number, paragraph in enumerate(paragraphs, start=1):
            text = paragraph.text
            identity = f"slide:{slide_number}:shape:{shape_number}:p:{paragraph_number}"
            extra: dict[str, JsonValue] = {"shape_name": shape.name}
            kind = NodeKind.HEADING if is_title and paragraph_number == 1 else NodeKind.PARAGRAPH
            if is_title and paragraph_number == 1:
                extra["heading_level"] = 1
            nodes.append(
                CanonicalNode(
                    node_id=canonical_node_id(document_id, identity),
                    kind=kind,
                    structural_identity=identity,
                    parent_id=slide_id,
                    child_order=shape_number,
                    text=text,
                    locator=SourceLocator(
                        path=(
                            "document",
                            "slide",
                            str(slide_number),
                            "shape",
                            str(shape_number),
                            "paragraph",
                            str(paragraph_number),
                        ),
                        extra_fields={
                            "slide": slide_number,
                            "shape": shape_number,
                        },
                    ),
                    extra_fields=extra,
                )
            )


def _emit_notes(
    document_id: uuid.UUID,
    nodes: list[CanonicalNode],
    slide_id: uuid.UUID,
    slide_number: int,
    slide: Any,
) -> None:
    if not slide.has_notes_slide:
        return
    notes_slide = slide.notes_slide
    for paragraph_number, paragraph in enumerate(
        notes_slide.notes_text_frame.paragraphs, start=1
    ):
        text = paragraph.text
        if not text:
            continue
        identity = f"slide:{slide_number}:notes:p:{paragraph_number}"
        nodes.append(
            CanonicalNode(
                node_id=canonical_node_id(document_id, identity),
                kind=NodeKind.PARAGRAPH,
                structural_identity=identity,
                parent_id=slide_id,
                child_order=_MAX_SHAPES + paragraph_number,
                text=text,
                locator=SourceLocator(
                    path=(
                        "document",
                        "slide",
                        str(slide_number),
                        "notes",
                        "paragraph",
                        str(paragraph_number),
                    ),
                    extra_fields={"slide": slide_number, "notes": True},
                ),
            )
        )


def _is_title_placeholder(shape: Any) -> bool:
    try:
        return shape.is_placeholder and shape.placeholder_format.idx in (0, 1)
    except Exception:
        return False
