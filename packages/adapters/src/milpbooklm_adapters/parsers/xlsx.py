"""
Deterministic XLSX canonical parser (openpyxl read-only, data-only).

Values and formulas are stored distinctly: cell text is the cached/displayed
value from a ``data_only`` load, while a formula (from a separate read-only
load) lands in the cell's ``formula`` extra field. Nothing is ever evaluated
— no calculation, no LibreOffice recalculation, no external-link resolution
(workbooks with external links are policy-rejected up front). Locator
guarantees: sheet nodes by name and cell nodes by A1 reference with
``sheet``/``row``/``col`` locator extras.
"""

from __future__ import annotations

import io
import uuid
import zipfile
from importlib.metadata import version
from typing import Any, Final

from milpbooklm_contracts.canonical_document import (
    CanonicalDocument,
    CanonicalNode,
    JsonValue,
    NodeKind,
    ParserDescriptor,
    SourceLocator,
    canonical_document_id,
    canonical_node_id,
)

from .office import (
    OfficeCorruptError,
    OfficePolicyError,
    OfficeTooLargeError,
    assert_no_macros,
    open_bounded_archive,
)

_PROFILE: Final = "xlsx-values-v1"
_MAX_SHEETS: Final = 200
_MAX_ROWS: Final = 100_000
_MAX_COLUMNS: Final = 1_024


def parse_xlsx(source_version_id: uuid.UUID, data: bytes) -> CanonicalDocument:
    """Extract cached cell values and distinct formulas without evaluation."""
    archive = open_bounded_archive(data)
    assert_no_macros(archive, "xl/")
    _assert_no_external_links(archive)
    archive.close()
    import openpyxl  # noqa: PLC0415 - parser imports load only inside the child

    try:
        values = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise OfficeCorruptError("workbook cannot be opened read-only") from exc
    if len(values.sheetnames) > _MAX_SHEETS:
        raise OfficeTooLargeError(f"workbook exceeds {_MAX_SHEETS} sheet limit")
    sheet_names = list(values.sheetnames)
    parser = ParserDescriptor(
        identity="milpbooklm.xlsx",
        version="1",
        profile=_PROFILE,
        tool_versions=(f"openpyxl={version('openpyxl')}",),
        extra_fields={"formula_policy": "stored_not_evaluated", "recalculation": "none"},
    )
    document_id = canonical_document_id(source_version_id, parser)
    root_id = canonical_node_id(document_id, "document")
    nodes = [
        CanonicalNode(
            node_id=root_id,
            kind=NodeKind.DOCUMENT,
            structural_identity="document",
            parent_id=None,
            child_order=0,
            text=None,
            locator=SourceLocator(path=("document",)),
        )
    ]
    for sheet_number, name in enumerate(sheet_names, start=1):
        sheet_identity = f"sheet:{sheet_number}"
        sheet_id = canonical_node_id(document_id, sheet_identity)
        nodes.append(
            CanonicalNode(
                node_id=sheet_id,
                kind=NodeKind.SHEET,
                structural_identity=sheet_identity,
                parent_id=root_id,
                child_order=sheet_number,
                text=None,
                locator=SourceLocator(
                    path=("document", "sheet", name),
                    extra_fields={"sheet": name, "sheet_index": sheet_number},
                ),
                extra_fields={"sheet_name": name},
            )
        )
        _emit_sheet(
            document_id, nodes, sheet_id, name, values=values[name], formulas=formulas[name]
        )
    values.close()
    formulas.close()
    return CanonicalDocument(
        document_id=document_id,
        source_version_id=source_version_id,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        parser=parser,
        languages=(),
        metadata={"sheets": sheet_names},
        root_node_ids=(root_id,),
        nodes=tuple(nodes),
    )


def _emit_sheet(
    document_id: uuid.UUID,
    nodes: list[CanonicalNode],
    sheet_id: uuid.UUID,
    name: str,
    *,
    values: Any,
    formulas: Any,
) -> None:
    table_identity = f"sheet:{name}:table:1"
    table_id = canonical_node_id(document_id, table_identity)
    nodes.append(
        CanonicalNode(
            node_id=table_id,
            kind=NodeKind.TABLE,
            structural_identity=table_identity,
            parent_id=sheet_id,
            child_order=1,
            text=None,
            locator=SourceLocator(
                path=("document", "sheet", name, "table", "1"),
                extra_fields={"sheet": name},
            ),
        )
    )
    for row_number, (value_row, formula_row) in enumerate(
        zip(values.iter_rows(), formulas.iter_rows(), strict=True), start=1
    ):
        if row_number > _MAX_ROWS:
            raise OfficeTooLargeError(f"sheet {name} exceeds {_MAX_ROWS} row limit")
        row_identity = f"{table_identity}:row:{row_number}"
        row_id = canonical_node_id(document_id, row_identity)
        nodes.append(
            CanonicalNode(
                node_id=row_id,
                kind=NodeKind.ROW,
                structural_identity=row_identity,
                parent_id=table_id,
                child_order=row_number,
                text=None,
                locator=SourceLocator(
                    path=("document", "sheet", name, "row", str(row_number)),
                    extra_fields={"sheet": name, "row": row_number},
                ),
            )
        )
        for column_number, (cell, formula_cell) in enumerate(
            zip(value_row, formula_row, strict=True), start=1
        ):
            if column_number > _MAX_COLUMNS:
                raise OfficeTooLargeError(f"sheet {name} exceeds {_MAX_COLUMNS} column limit")
            formula = _formula_text(formula_cell.value)
            extra: dict[str, JsonValue] = {}
            if formula is not None:
                extra["formula"] = formula
            nodes.append(
                CanonicalNode(
                    node_id=canonical_node_id(
                        document_id, f"{row_identity}:cell:{cell.coordinate}"
                    ),
                    kind=NodeKind.CELL,
                    structural_identity=f"{row_identity}:cell:{cell.coordinate}",
                    parent_id=row_id,
                    child_order=column_number,
                    text=_cell_text(cell.value),
                    locator=SourceLocator(
                        path=("document", "sheet", name, "cell", cell.coordinate),
                        extra_fields={
                            "sheet": name,
                            "row": cell.row,
                            "col": cell.column,
                            "cell_ref": cell.coordinate,
                        },
                    ),
                    extra_fields=extra,
                )
            )


def _formula_text(value: object) -> str | None:
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _assert_no_external_links(archive: zipfile.ZipFile) -> None:
    if any(info.filename.startswith("xl/externalLinks/") for info in archive.infolist()):
        raise OfficePolicyError("workbook contains external links")
